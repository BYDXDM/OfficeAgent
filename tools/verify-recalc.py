# -*- coding: utf-8 -*-
# LO 实算复验：读 lo_recalc/*.xlsx 的缓存值（data_only=True），逐值核对。
# 口径：社保=基本×(8%+2%+0.5%)，公积金=基本×12%，起征点 5000，月度税率表。
import os, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl

D = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "lo_recalc")
D = os.path.abspath(D)

fails = []
def chk(name, got, want, tol=0.005):
    if want is None:
        ok = (got is None)
    elif isinstance(want, str) or isinstance(want, list):
        ok = (got == want)
    else:
        ok = got is not None and abs(float(got) - float(want)) <= tol
    print(("  ok    " if ok else "  FAIL  ") + name + "  got=" + repr(got) + " want=" + repr(want))
    if not ok: fails.append(name)

def load(fn):
    return openpyxl.load_workbook(os.path.join(D, fn), data_only=True)

# ---- 工资表 ----
wb = load("工资表-2026年1月.xlsx")
print("工资表 sheets=", wb.sheetnames, "active=", wb.active.title)
chk("工资表页序", wb.sheetnames, ["工资表", "参数", "说明"])
ws = wb["工资表"]
chk("G2 应发", ws["G2"].value, 9500)
chk("H2 社保", ws["H2"].value, 840)
chk("J2 应纳税所得", ws["J2"].value, 2700)
chk("K2 个税", ws["K2"].value, 81)
chk("L2 实发", ws["L2"].value, 7619)
chk("K55 个税合计", ws["K55"].value, 94.5)
chk("L55 实发合计", ws["L55"].value, 17230.5)

# ---- 流水账 ----
wb = load("流水账.xlsx")
print("流水账 sheets=", wb.sheetnames, "active=", wb.active.title)
chk("流水账页序", wb.sheetnames, ["流水", "汇总", "参数", "说明"])
ws = wb["流水"]
chk("F2 余额", ws["F2"].value, 4000)
chk("F3 余额", ws["F3"].value, 3880)
chk("速览 I2 收入", ws["I2"].value, 3000)
chk("速览 I3 支出", ws["I3"].value, 120)
chk("速览 I4 结余", ws["I4"].value, 3880)
ws = wb["汇总"]
chk("汇总 B2 收入", ws["B2"].value, 3000)
chk("汇总 B3 支出", ws["B3"].value, 120)
chk("汇总 B4 结余", ws["B4"].value, 3880)

# ---- 增值税台账 ----
wb = load("增值税台账.xlsx")
ws = wb["台账"]
chk("台账 F2 税额", ws["F2"].value, 14690)
chk("台账 F3 税额", ws["F3"].value, 7345)
ws = wb["汇总"]
for i, w in enumerate([113000, 14690, 56500, 7345, 7345]):
    chk("汇总 B%d" % (i + 2), ws.cell(row=i + 2, column=2).value, w)

# ---- 自由表（本次修复核心）----
wb = load("自由表.xlsx")
print("自由表 sheets=", wb.sheetnames)
ws = wb["明细"]
chk("明细 D2 金额", ws["D2"].value, 250)
chk("明细 D3 金额", ws["D3"].value, 150)
chk("明细 D4 空行", ws["D4"].value, 0)   # 自由模式 formulaCols 无 IF 守卫，空行算 0（既有设计，非缺陷）
ws = wb["汇总"]
print("汇总页全部单元格：")
for r in range(1, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 2)]
    print("   row", r, vals)
chk("汇总 A2", ws["A2"].value, "办公用品")
chk("汇总 B2", ws["B2"].value, 250)
chk("汇总 A3", ws["A3"].value, "交通费")
chk("汇总 B3", ws["B3"].value, 150)
chk("汇总总计=400（不双计）", ws["B4"].value, 400)
chk("汇总无伪分组「合计」（A4 非『合计』字样冲突计数）",
    sum(1 for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "合计"), 1)

print()
print("RECALC ALL PASS" if not fails else ("RECALC FAILED=" + str(len(fails)) + " " + ",".join(fails)))
sys.exit(0 if not fails else 2)
